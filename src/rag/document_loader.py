"""
Belge okuma modülü.

Farklı dosya formatlarından düz metin çıkarır.
Her format için ayrı bir okuyucu fonksiyon (SOLID - Single Responsibility).

Desteklenen formatlar:
- PDF   → pymupdf (fitz)
- DOCX  → python-docx
- XLSX  → openpyxl
- CSV   → dahili csv modülü
- TXT, MD, kod dosyaları → doğrudan okuma

Neden pymupdf?
- Saf Python PDF kütüphanelerinden (PyPDF2) çok daha hızlı ve doğru metin çıkarır.
- Tablolar ve karmaşık düzenlerde bile iyi sonuç verir.
"""

import csv
from pathlib import Path

import fitz  # pymupdf
import openpyxl
from docx import Document as DocxDocument

# Doğrudan metin olarak okunabilen uzantılar
_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst",
    ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt",
    ".html", ".css", ".json", ".yaml", ".yml", ".xml", ".toml",
    ".sh", ".bat", ".ps1", ".sql",
}


def load_document(file_path: str | Path) -> str:
    """
    Bir belgeyi okur ve düz metin döner.

    Args:
        file_path: Okunacak dosyanın yolu

    Returns:
        Belgeden çıkarılan düz metin

    Raises:
        ValueError: Desteklenmeyen dosya formatı
        FileNotFoundError: Dosya bulunamadı
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Dosya bulunamadi: {path}")

    ext = path.suffix.lower()

    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext == ".xlsx":
        return _load_xlsx(path)
    if ext == ".csv":
        return _load_csv(path)
    if ext in _TEXT_EXTENSIONS:
        return _load_text(path)

    raise ValueError(f"Desteklenmeyen dosya formati: {ext}")


def is_supported(file_path: str | Path) -> bool:
    """Dosya formatının desteklenip desteklenmediğini kontrol eder."""
    ext = Path(file_path).suffix.lower()
    return ext in _TEXT_EXTENSIONS or ext in {".pdf", ".docx", ".xlsx", ".csv"}


def _load_pdf(path: Path) -> str:
    """PDF'den metin çıkarır (sayfa sayfa)."""
    parts: list[str] = []
    with fitz.open(path) as doc:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                parts.append(f"[Sayfa {page_num}]\n{text}")
    return "\n\n".join(parts)


def _load_docx(path: Path) -> str:
    """DOCX'ten paragraf metinlerini çıkarır."""
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(parts)


def _load_xlsx(path: Path) -> str:
    """Excel'den hücre verilerini satır satır çıkarır."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"[Sayfa: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _load_csv(path: Path) -> str:
    """CSV'yi satır satır okur."""
    parts: list[str] = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            if any(cell.strip() for cell in row):
                parts.append(" | ".join(row))
    return "\n".join(parts)


def _load_text(path: Path) -> str:
    """Düz metin dosyalarını okur."""
    return path.read_text(encoding="utf-8", errors="replace")
